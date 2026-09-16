import os
import openpyxl
from openpyxl.styles import Alignment

from .styles import (
    BANNER_FILL, BANNER_FONT, HEADER_FILL, HEADER_FONT,
    PART_FILL, PART_FONT, DATA_FONT, SUBTOTAL_FILL, SUBTOTAL_FONT,
    TOTAL_FILL, TOTAL_FONT, BORDER_THIN, BORDER_DOUBLE_BOTTOM,
    PARTS_CONFIG_AGENCY, _setup_portrait_page, _get_date_range_str,
    _apply_value_based_widths, _style_banner_row
)
from .agency_sheets import (
    _create_agency_print_all_sheet,
    _create_agency_breakup_sheet,
    _create_raw_data_sheet
)


def export_agency_workbooks(processed_result, output_dir):
    """
    Exports individual <Agency>.xlsx workbooks for each agency.
    Sheet structure:
    1. PRINT_ALL: The all-in-one printable bill:
       - Page 1: Breakup report (no signatures, no dual action word)
       - Page Break
       - Page 2: 1PH details (DR, DC, RC with countable SL and merged & centered total counts; NO amount; spacings)
       - Page Break (1PH and 3PH on different pages)
       - Page 3: 3PH details (DR, DC, RC with countable SL and merged & centered total counts; NO amount; spacings)
    2. BREAKUP REPORT: Standalone 1-page Breakup summary.
    3. Phase 1: Standalone Phase 1 detailed list with DR, DC, RC.
    4. Phase 3: Standalone Phase 3 detailed list with DR, DC, RC.
    5. RAW_DATA: Pure unformatted raw table data.
    """
    os.makedirs(output_dir, exist_ok=True)
    records = processed_result.get("records", [])
    date_range_str = _get_date_range_str(processed_result)

    agencies_map = {}
    for r in records:
        ag = r.get("agency") or "UNASSIGNED"
        if ag not in agencies_map:
            agencies_map[ag] = []
        agencies_map[ag].append(r)

    exported_files = []
    headers = ['SL', 'Consumer ID', 'Payment Type', 'Amount', 'Payment Date', 'DOC NUMBER', 'ZONE', 'Phase']

    sorted_agencies = sorted([a for a in agencies_map if a != "UNASSIGNED"])
    if "UNASSIGNED" in agencies_map:
        sorted_agencies.append("UNASSIGNED")

    for file_num, agency_name in enumerate(sorted_agencies, start=3):
        ag_records = agencies_map[agency_name]
        wb = openpyxl.Workbook()
        default_sheet = wb.active

        # Sheet 1: The All-In-One Printable Bill
        _create_agency_print_all_sheet(wb, agency_name, processed_result, ag_records)

        # Sheet 2: Standalone Breakup Report
        _create_agency_breakup_sheet(wb, agency_name, processed_result, ag_records)

        # Sheets 3 & 4: Standalone Phase 1 and Phase 3 sheets
        for phase_num in (1, 3):
            ws = wb.create_sheet(title=f"Phase {phase_num}")

            banner_text = f"AGENCY: {agency_name}   |   PHASE {phase_num}   |   DATE RANGE: {date_range_str}"
            ws.append([banner_text])
            _style_banner_row(ws, 1, len(headers), banner_text, fill=BANNER_FILL, font=BANNER_FONT, height=24)

            ws.append(headers)
            ws.row_dimensions[2].height = 24
            for col_idx, h in enumerate(headers, start=1):
                cell = ws.cell(2, col_idx)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = BORDER_THIN

            phase_records = [r for r in ag_records if r.get("phase") == phase_num]
            cur_part_idx = 1
            phase_total_count = 0

            for ptype, part_label in PARTS_CONFIG_AGENCY:
                part_records = [
                    r for r in phase_records
                    if str(r.get("payment_type") or "").strip().upper() == ptype
                ]
                if not part_records:
                    continue

                phase_total_count += len(part_records)

                # Part Banner (NO amount, Dual Action removed)
                banner_row = ws.max_row + 1
                part_title = f"PART {cur_part_idx}: {phase_num} PHASE - {part_label}   |   COUNT: {len(part_records)}"
                ws.append([part_title])
                ws.merge_cells(start_row=banner_row, start_column=1, end_row=banner_row, end_column=len(headers))
                ws.row_dimensions[banner_row].height = 21
                for c in range(1, len(headers) + 1):
                    cell = ws.cell(banner_row, c)
                    cell.fill = PART_FILL
                    cell.border = BORDER_THIN
                top_c = ws.cell(banner_row, 1)
                top_c.font = PART_FONT
                top_c.alignment = Alignment(horizontal="left", vertical="center", indent=1)

                for sl, r in enumerate(part_records, start=1):
                    try:
                        amt_val = float(r["amount"])
                    except (ValueError, TypeError):
                        amt_val = r["amount"]

                    row_vals = [
                        sl, r["consumer_id"], r["payment_type"], amt_val,
                        r["payment_date"], r["doc_number"], r["zone"], r["phase"]
                    ]
                    ws.append(row_vals)
                    cur_row = ws.max_row
                    ws.row_dimensions[cur_row].height = 19

                    for c in range(1, len(row_vals) + 1):
                        cell = ws.cell(cur_row, c)
                        cell.font = DATA_FONT
                        cell.border = BORDER_THIN
                        if c == 4 and isinstance(amt_val, (int, float)):
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                            cell.number_format = '#,##0.00'
                        else:
                            cell.alignment = Alignment(horizontal="center", vertical="center")

                # Subtotal row: NO TOTAL AMOUNT, MERGED AND CENTERED TOTAL COUNT
                subtotal_row = ws.max_row + 1
                ws.append([f"TOTAL {phase_num}PH {ptype} RECORDS: {len(part_records)}"] + [""] * (len(headers) - 1))
                ws.merge_cells(start_row=subtotal_row, start_column=1, end_row=subtotal_row, end_column=len(headers))
                ws.row_dimensions[subtotal_row].height = 20
                for c in range(1, len(headers) + 1):
                    cell = ws.cell(subtotal_row, c)
                    cell.fill = SUBTOTAL_FILL
                    cell.border = BORDER_DOUBLE_BOTTOM
                mid_c = ws.cell(subtotal_row, 1)
                mid_c.font = SUBTOTAL_FONT
                mid_c.alignment = Alignment(horizontal="center", vertical="center")

                # Spacing after parts
                ws.append([""] * len(headers))
                ws.row_dimensions[ws.max_row].height = 10

                cur_part_idx += 1

            if phase_total_count > 0:
                tot_row = ws.max_row + 1
                ws.append([f"TOTAL {phase_num} PHASE ({phase_num}PH) RECORDS: {phase_total_count}"] + [""] * (len(headers) - 1))
                ws.merge_cells(start_row=tot_row, start_column=1, end_row=tot_row, end_column=len(headers))
                ws.row_dimensions[tot_row].height = 22
                for c in range(1, len(headers) + 1):
                    cell = ws.cell(tot_row, c)
                    cell.fill = TOTAL_FILL
                    cell.border = BORDER_DOUBLE_BOTTOM
                ws.cell(tot_row, 1).font = TOTAL_FONT
                ws.cell(tot_row, 1).alignment = Alignment(horizontal="center", vertical="center")
            else:
                empty_row = ws.max_row + 1
                ws.append(["", f"No Phase {phase_num} records for {agency_name}", "", "", "", "", "", ""])
                ws.merge_cells(start_row=empty_row, start_column=1, end_row=empty_row, end_column=len(headers))
                for c in range(1, len(headers) + 1):
                    cell = ws.cell(empty_row, c)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.font = DATA_FONT

            _setup_portrait_page(ws, title_rows='1:2', fit_to_height=0)
            _apply_value_based_widths(ws, headers, data_start_row=3)

        # Sheet 5: Pure Raw Data Sheet (Zero formatting)
        _create_raw_data_sheet(wb, headers, ag_records, is_master=False)

        wb.remove(default_sheet)

        safe_name = "".join(c for c in agency_name if c.isalnum() or c in (' ', '_', '-')).strip()
        file_path = os.path.join(output_dir, f"{file_num:02d}_{safe_name}.xlsx")
        wb.save(file_path)
        exported_files.append(file_path)

    return exported_files
