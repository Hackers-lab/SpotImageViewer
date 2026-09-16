import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# --- LIGHT THEME PALETTE (Toner-Friendly, Crisp, High-Contrast) ---
BANNER_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")  # Slate 200
BANNER_FONT = Font(name="Calibri", size=11, bold=True, color="0F172A")  # Slate 900

BANNER_SUB_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")  # Slate 100
BANNER_SUB_FONT = Font(name="Calibri", size=10, bold=True, color="1E293B")  # Slate 800

HEADER_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")  # Slate 200
HEADER_FONT = Font(name="Calibri", size=10, bold=True, color="0F172A")

SUBHEADER_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")  # Slate 100
SUBHEADER_FONT = Font(name="Calibri", size=9, bold=True, color="1E293B")

PART_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")  # Slate 200
PART_FONT = Font(name="Calibri", size=10, bold=True, color="0F172A")

DATA_FONT = Font(name="Calibri", size=10, color="0F172A")

SUBTOTAL_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")  # Slate 100
SUBTOTAL_FONT = Font(name="Calibri", size=10, bold=True, color="0F172A")

TOTAL_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")  # Slate 200
TOTAL_FONT = Font(name="Calibri", size=10, bold=True, color="0F172A")

BORDER_THIN = Border(
    left=Side(style="thin", color="94A3B8"),
    right=Side(style="thin", color="94A3B8"),
    top=Side(style="thin", color="94A3B8"),
    bottom=Side(style="thin", color="94A3B8")
)

BORDER_DOUBLE_BOTTOM = Border(
    left=Side(style="thin", color="94A3B8"),
    right=Side(style="thin", color="94A3B8"),
    top=Side(style="thin", color="94A3B8"),
    bottom=Side(style="double", color="0F172A")
)

DEFAULT_COL_WIDTHS_PORTRAIT = {
    'SL': 5,
    'Consumer ID': 12,
    'Payment Type': 7,
    'Amount': 10,
    'Payment Date': 11,
    'DOC NUMBER': 13,
    'ZONE': 10,
    'Agency': 14,
    'Phase': 6
}

PARTS_CONFIG_AGENCY = [
    ("DR", "DR (Disconnection & Reconnection)"),
    ("DC", "DC (Disconnection)"),
    ("RC", "RC (Reconnection)"),
]


def _setup_portrait_page(ws, title_rows='1:2', fit_to_height=0):
    """
    Configures Portrait A4 printing with equal left/right margins,
    horizontal centering so data fits equally from both sides,
    repeating title rows across pages, and page numbering.
    """
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.4
    ws.page_margins.header = 0.2
    ws.page_margins.footer = 0.2
    ws.print_options.horizontalCentered = True
    ws.print_options.verticalCentered = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = fit_to_height
    if title_rows:
        ws.print_title_rows = title_rows
    ws.oddFooter.right.text = "Page &P of &N"
    ws.evenFooter.right.text = "Page &P of &N"
    ws.oddFooter.left.text = ""
    ws.evenFooter.left.text = ""
    ws.print_options.gridLines = True
    if ws.views and ws.views.sheetView:
        ws.views.sheetView[0].showGridLines = True


def _get_date_range_str(processed_result):
    dr = processed_result.get("date_range", "")
    if dr:
        return dr
    dates = []
    for r in processed_result.get("records", []):
        pd = r.get("payment_date")
        if pd:
            dates.append(str(pd))
    if dates:
        return f"{min(dates)} to {max(dates)}"
    return "N/A"


def _apply_value_based_widths(ws, headers, data_start_row=3):
    """
    Adjusts column widths based strictly on data rows (ignoring banners and merged headers).
    Headers are styled with wrap_text=True so table columns stay compact for Portrait printing.
    """
    max_lens = {i: 0 for i in range(1, len(headers) + 1)}

    for row in ws.iter_rows(min_row=data_start_row, values_only=True):
        for col_idx, val in enumerate(row, start=1):
            if val is not None and col_idx <= len(headers):
                if col_idx == 1 and not isinstance(val, int):
                    continue
                if col_idx == 2 and isinstance(val, str) and any(w in val for w in ("PART", "TOTAL", "GRAND")):
                    continue
                s = str(val).strip()
                if '\n' in s:
                    s = max(s.split('\n'), key=len)
                max_lens[col_idx] = max(max_lens[col_idx], len(s))

    for col_idx, h_text in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        def_w = DEFAULT_COL_WIDTHS_PORTRAIT.get(h_text, 10)
        final_w = max(max_lens[col_idx] + 3, def_w)
        ws.column_dimensions[col_letter].width = final_w


def _style_banner_row(ws, start_col, end_col, text, fill=BANNER_FILL, font=BANNER_FONT, height=24):
    """Creates a light unified header banner across columns start_col to end_col."""
    cur_row = ws.max_row if ws.max_row > 0 else 1
    ws.merge_cells(start_row=cur_row, start_column=start_col, end_row=cur_row, end_column=end_col)
    ws.row_dimensions[cur_row].height = height
    for c in range(start_col, end_col + 1):
        cell = ws.cell(cur_row, c)
        cell.fill = fill
        cell.border = BORDER_THIN
    top_left = ws.cell(cur_row, start_col)
    top_left.value = text
    top_left.font = font
    top_left.alignment = Alignment(horizontal="center", vertical="center")


def _get_agency_metrics(agency_name, processed_result, ag_records):
    """Extracts or computes counts, rates, and amounts for an agency."""
    summary_item = None
    for s in processed_result.get("summary", []):
        if s.get("agency") == agency_name:
            summary_item = s
            break

    if summary_item:
        c_1_dr = summary_item.get("c_1ph_dr", 0)
        r_1_dr = summary_item.get("r_1ph_dr", 130.0)
        c_1_dc = summary_item.get("c_1ph_dc", 0)
        r_1_dc = summary_item.get("r_1ph_dc", 65.0)
        c_1_rc = summary_item.get("c_1ph_rc", 0)
        r_1_rc = summary_item.get("r_1ph_rc", 65.0)

        c_3_dr = summary_item.get("c_3ph_dr", 0)
        r_3_dr = summary_item.get("r_3ph_dr", 176.0)
        c_3_dc = summary_item.get("c_3ph_dc", 0)
        r_3_dc = summary_item.get("r_3ph_dc", 88.0)
        c_3_rc = summary_item.get("c_3ph_rc", 0)
        r_3_rc = summary_item.get("r_3ph_rc", 88.0)
    else:
        c_1_dr = c_1_dc = c_1_rc = 0
        c_3_dr = c_3_dc = c_3_rc = 0
        for r in ag_records:
            phase = r.get("phase")
            ptype = str(r.get("payment_type") or "").strip().upper()
            if phase == 1:
                if ptype == "DR": c_1_dr += 1
                elif ptype == "DC": c_1_dc += 1
                elif ptype == "RC": c_1_rc += 1
            elif phase == 3:
                if ptype == "DR": c_3_dr += 1
                elif ptype == "DC": c_3_dc += 1
                elif ptype == "RC": c_3_rc += 1

        r_1_dr, r_1_dc, r_1_rc = 130.0, 65.0, 65.0
        r_3_dr, r_3_dc, r_3_rc = 176.0, 88.0, 88.0

    amt_1_dr = c_1_dr * r_1_dr
    amt_1_dc = c_1_dc * r_1_dc
    amt_1_rc = c_1_rc * r_1_rc
    sub_1_count = c_1_dr + c_1_dc + c_1_rc
    sub_1_amt = amt_1_dr + amt_1_dc + amt_1_rc

    amt_3_dr = c_3_dr * r_3_dr
    amt_3_dc = c_3_dc * r_3_dc
    amt_3_rc = c_3_rc * r_3_rc
    sub_3_count = c_3_dr + c_3_dc + c_3_rc
    sub_3_amt = amt_3_dr + amt_3_dc + amt_3_rc

    grand_count = sub_1_count + sub_3_count
    grand_amt = sub_1_amt + sub_3_amt

    return {
        "c_1_dr": c_1_dr, "r_1_dr": r_1_dr, "amt_1_dr": amt_1_dr,
        "c_1_dc": c_1_dc, "r_1_dc": r_1_dc, "amt_1_dc": amt_1_dc,
        "c_1_rc": c_1_rc, "r_1_rc": r_1_rc, "amt_1_rc": amt_1_rc,
        "sub_1_count": sub_1_count, "sub_1_amt": sub_1_amt,
        "c_3_dr": c_3_dr, "r_3_dr": r_3_dr, "amt_3_dr": amt_3_dr,
        "c_3_dc": c_3_dc, "r_3_dc": r_3_dc, "amt_3_dc": amt_3_dc,
        "c_3_rc": c_3_rc, "r_3_rc": r_3_rc, "amt_3_rc": amt_3_rc,
        "sub_3_count": sub_3_count, "sub_3_amt": sub_3_amt,
        "grand_count": grand_count, "grand_amt": grand_amt
    }
