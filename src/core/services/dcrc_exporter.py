import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break


# --- LIGHT THEME PALETTE (Toner-Friendly, Crisp, High-Contrast) ---
BANNER_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")  # Slate 200
BANNER_FONT = Font(name="Calibri", size=11, bold=True, color="0F172A")  # Slate 900

BANNER_SUB_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")  # Slate 100
BANNER_SUB_FONT = Font(name="Calibri", size=10, bold=True, color="1E293B")  # Slate 800

HEADER_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")  # Slate 200
HEADER_FONT = Font(name="Calibri", size=10, bold=True, color="0F172A")

SUBHEADER_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")  # Slate 100
SUBHEADER_FONT = Font(name="Calibri", size=9, bold=True, color="1E293B")

PART_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")  # Slate 200
PART_FONT = Font(name="Calibri", size=10, bold=True, color="0F172A")

DATA_FONT = Font(name="Calibri", size=10, color="0F172A")

SUBTOTAL_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")  # Slate 100
SUBTOTAL_FONT = Font(name="Calibri", size=10, bold=True, color="0F172A")

TOTAL_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")  # Slate 200
TOTAL_FONT = Font(name="Calibri", size=10, bold=True, color="0F172A")

BORDER_THIN = Border(
    left=Side(style="thin", color="94A3B8"),
    right=Side(style="thin", color="94A3B8"),
    top=Side(style="thin", color="94A3B8"),
    bottom=Side(style="thin", color="94A3B8")
)

BORDER_DOUBLE_BOTTOM = Border(
    left=Side(style="thin", color="94A3B8"),
    right=Side(style="thin", color="94A3B8"),
    top=Side(style="thin", color="94A3B8"),
    bottom=Side(style="double", color="0F172A")
)

DEFAULT_COL_WIDTHS_PORTRAIT = {
    'SL': 5,
    'Consumer ID': 12,
    'Payment Type': 7,
    'Amount': 10,
    'Payment Date': 11,
    'DOC NUMBER': 13,
    'ZONE': 10,
    'Agency': 14,
    'Phase': 6
}

PARTS_CONFIG_AGENCY = [
    ("DR", "DR (Disconnection & Reconnection)"),
    ("DC", "DC (Disconnection)"),
    ("RC", "RC (Reconnection)"),
]


def _setup_portrait_page(ws, title_rows='1:2', fit_to_height=0):
    """
    Configures Portrait A4 printing with equal left/right margins,
    horizontal centering so data fits equally from both sides,
    repeating title rows across pages, and page numbering.
    """
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.4
    ws.page_margins.header = 0.2
    ws.page_margins.footer = 0.2
    ws.print_options.horizontalCentered = True
    ws.print_options.verticalCentered = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = fit_to_height
    if title_rows:
        ws.print_title_rows = title_rows
    ws.oddFooter.right.text = "Page &P of &N"
    ws.evenFooter.right.text = "Page &P of &N"
    ws.oddFooter.left.text = ""
    ws.evenFooter.left.text = ""
    ws.print_options.gridLines = True
    if ws.views and ws.views.sheetView:
        ws.views.sheetView[0].showGridLines = True


def _get_date_range_str(processed_result):
    dr = processed_result.get("date_range", "")
    if dr:
        return dr
    dates = []
    for r in processed_result.get("records", []):
        pd = r.get("payment_date")
        if pd:
            dates.append(str(pd))
    if dates:
        return f"{min(dates)} to {max(dates)}"
    return "N/A"


def _apply_value_based_widths(ws, headers, data_start_row=3):
    """
    Adjusts column widths based strictly on data rows (ignoring banners and merged headers).
    Headers are styled with wrap_text=True so table columns stay compact for Portrait printing.
    """
    max_lens = {i: 0 for i in range(1, len(headers) + 1)}
    
    for row in ws.iter_rows(min_row=data_start_row, values_only=True):
        for col_idx, val in enumerate(row, start=1):
            if val is not None and col_idx <= len(headers):
                if col_idx == 1 and not isinstance(val, int):
                    continue
                if col_idx == 2 and isinstance(val, str) and any(w in val for w in ("PART", "TOTAL", "GRAND")):
                    continue
                s = str(val).strip()
                if '\n' in s:
                    s = max(s.split('\n'), key=len)
                max_lens[col_idx] = max(max_lens[col_idx], len(s))
                
    for col_idx, h_text in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        def_w = DEFAULT_COL_WIDTHS_PORTRAIT.get(h_text, 10)
        final_w = max(max_lens[col_idx] + 3, def_w)
        ws.column_dimensions[col_letter].width = final_w


def _style_banner_row(ws, start_col, end_col, text, fill=BANNER_FILL, font=BANNER_FONT, height=24):
    """Creates a light unified header banner across columns start_col to end_col."""
    cur_row = ws.max_row if ws.max_row > 0 else 1
    ws.merge_cells(start_row=cur_row, start_column=start_col, end_row=cur_row, end_column=end_col)
    ws.row_dimensions[cur_row].height = height
    for c in range(start_col, end_col + 1):
        cell = ws.cell(cur_row, c)
        cell.fill = fill
        cell.border = BORDER_THIN
    top_left = ws.cell(cur_row, start_col)
    top_left.value = text
    top_left.font = font
    top_left.alignment = Alignment(horizontal="center", vertical="center")


def _get_agency_metrics(agency_name, processed_result, ag_records):
    """Extracts or computes counts, rates, and amounts for an agency."""
    summary_item = None
    for s in processed_result.get("summary", []):
        if s.get("agency") == agency_name:
            summary_item = s
            break

    if summary_item:
        c_1_dr = summary_item.get("c_1ph_dr", 0)
        r_1_dr = summary_item.get("r_1ph_dr", 130.0)
        c_1_dc = summary_item.get("c_1ph_dc", 0)
        r_1_dc = summary_item.get("r_1ph_dc", 65.0)
        c_1_rc = summary_item.get("c_1ph_rc", 0)
        r_1_rc = summary_item.get("r_1ph_rc", 65.0)

        c_3_dr = summary_item.get("c_3ph_dr", 0)
        r_3_dr = summary_item.get("r_3ph_dr", 176.0)
        c_3_dc = summary_item.get("c_3ph_dc", 0)
        r_3_dc = summary_item.get("r_3ph_dc", 88.0)
        c_3_rc = summary_item.get("c_3ph_rc", 0)
        r_3_rc = summary_item.get("r_3ph_rc", 88.0)
    else:
        c_1_dr = c_1_dc = c_1_rc = 0
        c_3_dr = c_3_dc = c_3_rc = 0
        for r in ag_records:
            phase = r.get("phase")
            ptype = str(r.get("payment_type") or "").strip().upper()
            if phase == 1:
                if ptype == "DR": c_1_dr += 1
                elif ptype == "DC": c_1_dc += 1
                elif ptype == "RC": c_1_rc += 1
            elif phase == 3:
                if ptype == "DR": c_3_dr += 1
                elif ptype == "DC": c_3_dc += 1
                elif ptype == "RC": c_3_rc += 1

        r_1_dr, r_1_dc, r_1_rc = 130.0, 65.0, 65.0
        r_3_dr, r_3_dc, r_3_rc = 176.0, 88.0, 88.0

    amt_1_dr = c_1_dr * r_1_dr
    amt_1_dc = c_1_dc * r_1_dc
    amt_1_rc = c_1_rc * r_1_rc
    sub_1_count = c_1_dr + c_1_dc + c_1_rc
    sub_1_amt = amt_1_dr + amt_1_dc + amt_1_rc

    amt_3_dr = c_3_dr * r_3_dr
    amt_3_dc = c_3_dc * r_3_dc
    amt_3_rc = c_3_rc * r_3_rc
    sub_3_count = c_3_dr + c_3_dc + c_3_rc
    sub_3_amt = amt_3_dr + amt_3_dc + amt_3_rc

    grand_count = sub_1_count + sub_3_count
    grand_amt = sub_1_amt + sub_3_amt

    return {
        "c_1_dr": c_1_dr, "r_1_dr": r_1_dr, "amt_1_dr": amt_1_dr,
        "c_1_dc": c_1_dc, "r_1_dc": r_1_dc, "amt_1_dc": amt_1_dc,
        "c_1_rc": c_1_rc, "r_1_rc": r_1_rc, "amt_1_rc": amt_1_rc,
        "sub_1_count": sub_1_count, "sub_1_amt": sub_1_amt,
        "c_3_dr": c_3_dr, "r_3_dr": r_3_dr, "amt_3_dr": amt_3_dr,
        "c_3_dc": c_3_dc, "r_3_dc": r_3_dc, "amt_3_dc": amt_3_dc,
        "c_3_rc": c_3_rc, "r_3_rc": r_3_rc, "amt_3_rc": amt_3_rc,
        "sub_3_count": sub_3_count, "sub_3_amt": sub_3_amt,
        "grand_count": grand_count, "grand_amt": grand_amt
    }


def _create_agency_breakup_sheet(wb, agency_name, processed_result, ag_records):
    """
    Creates the standalone Breakup Summary Report sheet for an agency.
    (Dual Action word removed, Signatures block removed as requested).
    Formatted to print on exactly 1 Portrait page with light headers and horizontal centering.
    """
    ws = wb.create_sheet(title="BREAKUP REPORT")
    date_range_str = _get_date_range_str(processed_result)
    m = _get_agency_metrics(agency_name, processed_result, ag_records)
    total_cols = 5

    # Row 1: Title Banner
    ws.append(["PURCHASE ORDER BILLING BREAKUP REPORT"])
    _style_banner_row(ws, 1, total_cols, "PURCHASE ORDER BILLING BREAKUP REPORT", fill=BANNER_FILL, font=BANNER_FONT, height=26)

    # Row 2: Subtitle Banner
    ws.append([f"AGENCY: {agency_name}   |   DATE RANGE: {date_range_str}"])
    _style_banner_row(ws, 1, total_cols, f"AGENCY: {agency_name}   |   DATE RANGE: {date_range_str}", fill=BANNER_SUB_FILL, font=BANNER_SUB_FONT, height=22)

    # Row 3: Spacer
    ws.append([""] * total_cols)
    ws.row_dimensions[3].height = 8

    # Row 4: Headers
    headers = ["SL", "Particulars / Work Item Description", "Records / Qty", "Rate (Rs.)", "Total Amount (Rs.)"]
    ws.append(headers)
    ws.row_dimensions[4].height = 24
    for c in range(1, total_cols + 1):
        cell = ws.cell(4, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_THIN

    # Section 1: Single Phase (Dual Action removed)
    sec1_row = ws.max_row + 1
    ws.append(["SINGLE PHASE (1PH) WORK ITEMS"])
    ws.merge_cells(start_row=sec1_row, start_column=1, end_row=sec1_row, end_column=total_cols)
    ws.row_dimensions[sec1_row].height = 20
    for c in range(1, total_cols + 1):
        cell = ws.cell(sec1_row, c)
        cell.fill = PART_FILL
        cell.border = BORDER_THIN
    top_cell = ws.cell(sec1_row, 1)
    top_cell.font = PART_FONT
    top_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    items_1ph = [
        (1, "1 Phase - DR (Disconnection & Reconnection)", m["c_1_dr"], m["r_1_dr"], m["amt_1_dr"]),
        (2, "1 Phase - DC (Disconnection)", m["c_1_dc"], m["r_1_dc"], m["amt_1_dc"]),
        (3, "1 Phase - RC (Reconnection)", m["c_1_rc"], m["r_1_rc"], m["amt_1_rc"])
    ]
    for sl, desc, cnt, rate, amt in items_1ph:
        ws.append([sl, desc, cnt, rate, amt])
        cur_r = ws.max_row
        ws.row_dimensions[cur_r].height = 20
        for c in range(1, total_cols + 1):
            cell = ws.cell(cur_r, c)
            cell.font = DATA_FONT
            cell.border = BORDER_THIN
            if c == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c == 2:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif c == 3:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c in (4, 5):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0.00'

    sub1_row = ws.max_row + 1
    ws.append(["", "SUBTOTAL SINGLE PHASE (1PH)", m["sub_1_count"], "", m["sub_1_amt"]])
    ws.row_dimensions[sub1_row].height = 21
    for c in range(1, total_cols + 1):
        cell = ws.cell(sub1_row, c)
        cell.font = SUBTOTAL_FONT
        cell.fill = SUBTOTAL_FILL
        cell.border = BORDER_DOUBLE_BOTTOM
        if c == 2:
            cell.alignment = Alignment(horizontal="left", vertical="center")
        elif c == 3:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif c == 5:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = '#,##0.00'

    # Section 2: Three Phase (Dual Action removed)
    sec3_row = ws.max_row + 1
    ws.append(["THREE PHASE (3PH) WORK ITEMS"])
    ws.merge_cells(start_row=sec3_row, start_column=1, end_row=sec3_row, end_column=total_cols)
    ws.row_dimensions[sec3_row].height = 20
    for c in range(1, total_cols + 1):
        cell = ws.cell(sec3_row, c)
        cell.fill = PART_FILL
        cell.border = BORDER_THIN
    top_cell = ws.cell(sec3_row, 1)
    top_cell.font = PART_FONT
    top_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    items_3ph = [
        (4, "3 Phase - DR (Disconnection & Reconnection)", m["c_3_dr"], m["r_3_dr"], m["amt_3_dr"]),
        (5, "3 Phase - DC (Disconnection)", m["c_3_dc"], m["r_3_dc"], m["amt_3_dc"]),
        (6, "3 Phase - RC (Reconnection)", m["c_3_rc"], m["r_3_rc"], m["amt_3_rc"])
    ]
    for sl, desc, cnt, rate, amt in items_3ph:
        ws.append([sl, desc, cnt, rate, amt])
        cur_r = ws.max_row
        ws.row_dimensions[cur_r].height = 20
        for c in range(1, total_cols + 1):
            cell = ws.cell(cur_r, c)
            cell.font = DATA_FONT
            cell.border = BORDER_THIN
            if c == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c == 2:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif c == 3:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c in (4, 5):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0.00'

    sub3_row = ws.max_row + 1
    ws.append(["", "SUBTOTAL THREE PHASE (3PH)", m["sub_3_count"], "", m["sub_3_amt"]])
    ws.row_dimensions[sub3_row].height = 21
    for c in range(1, total_cols + 1):
        cell = ws.cell(sub3_row, c)
        cell.font = SUBTOTAL_FONT
        cell.fill = SUBTOTAL_FILL
        cell.border = BORDER_DOUBLE_BOTTOM
        if c == 2:
            cell.alignment = Alignment(horizontal="left", vertical="center")
        elif c == 3:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif c == 5:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = '#,##0.00'

    # Grand Total
    gt_row = ws.max_row + 1
    ws.append(["", "GRAND TOTAL PAYABLE", m["grand_count"], "", m["grand_amt"]])
    ws.row_dimensions[gt_row].height = 24
    for c in range(1, total_cols + 1):
        cell = ws.cell(gt_row, c)
        cell.font = TOTAL_FONT
        cell.fill = TOTAL_FILL
        cell.border = BORDER_DOUBLE_BOTTOM
        if c == 2:
            cell.alignment = Alignment(horizontal="left", vertical="center")
        elif c == 3:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif c == 5:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = '#,##0.00'

    # Highlights Box
    sp_row = ws.max_row + 1
    ws.append([""] * total_cols)
    ws.row_dimensions[sp_row].height = 14

    box_start = ws.max_row + 1
    ws.append(["SUMMARY HIGHLIGHTS", "", "", "", ""])
    ws.merge_cells(start_row=box_start, start_column=1, end_row=box_start, end_column=total_cols)
    ws.cell(box_start, 1).font = TOTAL_FONT
    ws.cell(box_start, 1).alignment = Alignment(horizontal="left")

    ws.append(["1-Phase Total Records:", m["sub_1_count"], "", "1-Phase Payable Amount (Rs):", m["sub_1_amt"]])
    ws.append(["3-Phase Total Records:", m["sub_3_count"], "", "3-Phase Payable Amount (Rs):", m["sub_3_amt"]])
    ws.append(["Overall Total Records:", m["grand_count"], "", "Overall P.O. Payable Amount (Rs):", m["grand_amt"]])

    for r_idx in range(box_start + 1, box_start + 4):
        ws.row_dimensions[r_idx].height = 18
        ws.cell(r_idx, 1).font = DATA_FONT
        ws.cell(r_idx, 2).font = TOTAL_FONT
        ws.cell(r_idx, 2).alignment = Alignment(horizontal="left")
        ws.cell(r_idx, 4).font = DATA_FONT
        c5 = ws.cell(r_idx, 5)
        c5.font = TOTAL_FONT
        c5.alignment = Alignment(horizontal="right")
        c5.number_format = '#,##0.00'

    # Signatures block removed as requested: "THIS PAGE DONT NEED SIGN"

    breakup_widths = {'A': 6, 'B': 44, 'C': 15, 'D': 14, 'E': 18}
    for col_l, w in breakup_widths.items():
        ws.column_dimensions[col_l].width = w

    _setup_portrait_page(ws, title_rows='', fit_to_height=1)
    return ws


def _create_agency_print_all_sheet(wb, agency_name, processed_result, ag_records):
    """
    Creates the comprehensive 'PRINT_ALL' sheet where user can print directly at once:
    - Page 1 contains the complete Breakup Report (without signatures, dual action word removed)
    - Followed by explicit page break so 1PH and 3PH are printed on separate pages!
    - Consumer list has NO total amount in subtotal rows; total counts are merged and centered!
    - Spacings added after parts.
    - Configured for Portrait A4, horizontal centering, repeating table headers, and page numbers.
    """
    ws = wb.create_sheet(title="PRINT_ALL")
    date_range_str = _get_date_range_str(processed_result)
    m = _get_agency_metrics(agency_name, processed_result, ag_records)
    total_cols = 8  # Matched to consumer columns: SL, CID, Type, Amt, Date, Doc, Zone, Phase

    # 1. Title Banner
    ws.append(["PURCHASE ORDER BILLING BREAKUP & DETAILS REPORT"])
    _style_banner_row(ws, 1, total_cols, "PURCHASE ORDER BILLING BREAKUP & DETAILS REPORT", fill=BANNER_FILL, font=BANNER_FONT, height=24)

    # 2. Subtitle Banner
    ws.append([f"AGENCY: {agency_name}   |   DATE RANGE: {date_range_str}"])
    _style_banner_row(ws, 1, total_cols, f"AGENCY: {agency_name}   |   DATE RANGE: {date_range_str}", fill=BANNER_SUB_FILL, font=BANNER_SUB_FONT, height=22)

    # 3. Spacer
    ws.append([""] * total_cols)
    ws.row_dimensions[3].height = 6

    # 4. Breakup Table Header (Spans A to H)
    hdr_row = ws.max_row + 1
    ws.append(["SL", "Particulars / Work Item Description", "", "", "Records / Qty", "Rate (Rs.)", "Total Amount (Rs.)", ""])
    ws.merge_cells(start_row=hdr_row, start_column=2, end_row=hdr_row, end_column=4)
    ws.merge_cells(start_row=hdr_row, start_column=7, end_row=hdr_row, end_column=8)
    ws.row_dimensions[hdr_row].height = 22
    for c in range(1, total_cols + 1):
        cell = ws.cell(hdr_row, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_THIN

    # Breakup: Single Phase Section Header (Dual Action removed)
    bsec1_r = ws.max_row + 1
    ws.append(["SINGLE PHASE (1PH) WORK ITEMS"] + [""] * (total_cols - 1))
    ws.merge_cells(start_row=bsec1_r, start_column=1, end_row=bsec1_r, end_column=total_cols)
    ws.row_dimensions[bsec1_r].height = 20
    for c in range(1, total_cols + 1):
        cell = ws.cell(bsec1_r, c)
        cell.fill = PART_FILL
        cell.border = BORDER_THIN
    ws.cell(bsec1_r, 1).font = PART_FONT
    ws.cell(bsec1_r, 1).alignment = Alignment(horizontal="left", vertical="center", indent=1)

    items_1ph = [
        (1, "1 Phase - DR (Disconnection & Reconnection)", m["c_1_dr"], m["r_1_dr"], m["amt_1_dr"]),
        (2, "1 Phase - DC (Disconnection)", m["c_1_dc"], m["r_1_dc"], m["amt_1_dc"]),
        (3, "1 Phase - RC (Reconnection)", m["c_1_rc"], m["r_1_rc"], m["amt_1_rc"])
    ]
    for sl, desc, cnt, rate, amt in items_1ph:
        ws.append([sl, desc, "", "", cnt, rate, amt, ""])
        cur_r = ws.max_row
        ws.merge_cells(start_row=cur_r, start_column=2, end_row=cur_r, end_column=4)
        ws.merge_cells(start_row=cur_r, start_column=7, end_row=cur_r, end_column=8)
        ws.row_dimensions[cur_r].height = 19
        for c in range(1, total_cols + 1):
            cell = ws.cell(cur_r, c)
            cell.font = DATA_FONT
            cell.border = BORDER_THIN
            if c == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c == 2:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif c == 5:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c in (6, 7):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0.00'

    # Breakup: Subtotal 1PH
    bsub1_r = ws.max_row + 1
    ws.append(["", "SUBTOTAL SINGLE PHASE (1PH)", "", "", m["sub_1_count"], "", m["sub_1_amt"], ""])
    ws.merge_cells(start_row=bsub1_r, start_column=2, end_row=bsub1_r, end_column=4)
    ws.merge_cells(start_row=bsub1_r, start_column=7, end_row=bsub1_r, end_column=8)
    ws.row_dimensions[bsub1_r].height = 20
    for c in range(1, total_cols + 1):
        cell = ws.cell(bsub1_r, c)
        cell.font = SUBTOTAL_FONT
        cell.fill = SUBTOTAL_FILL
        cell.border = BORDER_DOUBLE_BOTTOM
        if c == 2:
            cell.alignment = Alignment(horizontal="left", vertical="center")
        elif c == 5:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif c == 7:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = '#,##0.00'

    # Breakup: Three Phase Section Header (Dual Action removed)
    bsec3_r = ws.max_row + 1
    ws.append(["THREE PHASE (3PH) WORK ITEMS"] + [""] * (total_cols - 1))
    ws.merge_cells(start_row=bsec3_r, start_column=1, end_row=bsec3_r, end_column=total_cols)
    ws.row_dimensions[bsec3_r].height = 20
    for c in range(1, total_cols + 1):
        cell = ws.cell(bsec3_r, c)
        cell.fill = PART_FILL
        cell.border = BORDER_THIN
    ws.cell(bsec3_r, 1).font = PART_FONT
    ws.cell(bsec3_r, 1).alignment = Alignment(horizontal="left", vertical="center", indent=1)

    items_3ph = [
        (4, "3 Phase - DR (Disconnection & Reconnection)", m["c_3_dr"], m["r_3_dr"], m["amt_3_dr"]),
        (5, "3 Phase - DC (Disconnection)", m["c_3_dc"], m["r_3_dc"], m["amt_3_dc"]),
        (6, "3 Phase - RC (Reconnection)", m["c_3_rc"], m["r_3_rc"], m["amt_3_rc"])
    ]
    for sl, desc, cnt, rate, amt in items_3ph:
        ws.append([sl, desc, "", "", cnt, rate, amt, ""])
        cur_r = ws.max_row
        ws.merge_cells(start_row=cur_r, start_column=2, end_row=cur_r, end_column=4)
        ws.merge_cells(start_row=cur_r, start_column=7, end_row=cur_r, end_column=8)
        ws.row_dimensions[cur_r].height = 19
        for c in range(1, total_cols + 1):
            cell = ws.cell(cur_r, c)
            cell.font = DATA_FONT
            cell.border = BORDER_THIN
            if c == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c == 2:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif c == 5:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c in (6, 7):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0.00'

    # Breakup: Subtotal 3PH
    bsub3_r = ws.max_row + 1
    ws.append(["", "SUBTOTAL THREE PHASE (3PH)", "", "", m["sub_3_count"], "", m["sub_3_amt"], ""])
    ws.merge_cells(start_row=bsub3_r, start_column=2, end_row=bsub3_r, end_column=4)
    ws.merge_cells(start_row=bsub3_r, start_column=7, end_row=bsub3_r, end_column=8)
    ws.row_dimensions[bsub3_r].height = 20
    for c in range(1, total_cols + 1):
        cell = ws.cell(bsub3_r, c)
        cell.font = SUBTOTAL_FONT
        cell.fill = SUBTOTAL_FILL
        cell.border = BORDER_DOUBLE_BOTTOM
        if c == 2:
            cell.alignment = Alignment(horizontal="left", vertical="center")
        elif c == 5:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif c == 7:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = '#,##0.00'

    # Breakup: Grand Total Row
    bgt_r = ws.max_row + 1
    ws.append(["", "GRAND TOTAL PAYABLE", "", "", m["grand_count"], "", m["grand_amt"], ""])
    ws.merge_cells(start_row=bgt_r, start_column=2, end_row=bgt_r, end_column=4)
    ws.merge_cells(start_row=bgt_r, start_column=7, end_row=bgt_r, end_column=8)
    ws.row_dimensions[bgt_r].height = 22
    for c in range(1, total_cols + 1):
        cell = ws.cell(bgt_r, c)
        cell.font = TOTAL_FONT
        cell.fill = TOTAL_FILL
        cell.border = BORDER_DOUBLE_BOTTOM
        if c == 2:
            cell.alignment = Alignment(horizontal="left", vertical="center")
        elif c == 5:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif c == 7:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = '#,##0.00'

    # Breakup: Summary Highlights
    sp_row = ws.max_row + 1
    ws.append([""] * total_cols)
    ws.row_dimensions[sp_row].height = 14

    box_start = ws.max_row + 1
    ws.append(["SUMMARY HIGHLIGHTS", "", "", "", "", "", "", ""])
    ws.merge_cells(start_row=box_start, start_column=1, end_row=box_start, end_column=total_cols)
    ws.cell(box_start, 1).font = TOTAL_FONT
    ws.cell(box_start, 1).alignment = Alignment(horizontal="left")

    ws.append(["1-Phase Total Records:", m["sub_1_count"], "", "", "1-Phase Payable Amount (Rs):", "", m["sub_1_amt"], ""])
    ws.merge_cells(start_row=box_start + 1, start_column=5, end_row=box_start + 1, end_column=6)
    ws.merge_cells(start_row=box_start + 1, start_column=7, end_row=box_start + 1, end_column=8)

    ws.append(["3-Phase Total Records:", m["sub_3_count"], "", "", "3-Phase Payable Amount (Rs):", "", m["sub_3_amt"], ""])
    ws.merge_cells(start_row=box_start + 2, start_column=5, end_row=box_start + 2, end_column=6)
    ws.merge_cells(start_row=box_start + 2, start_column=7, end_row=box_start + 2, end_column=8)

    ws.append(["Overall Total Records:", m["grand_count"], "", "", "Overall P.O. Payable Amount (Rs):", "", m["grand_amt"], ""])
    ws.merge_cells(start_row=box_start + 3, start_column=5, end_row=box_start + 3, end_column=6)
    ws.merge_cells(start_row=box_start + 3, start_column=7, end_row=box_start + 3, end_column=8)

    for r_idx in range(box_start + 1, box_start + 4):
        ws.row_dimensions[r_idx].height = 18
        ws.cell(r_idx, 1).font = DATA_FONT
        ws.cell(r_idx, 2).font = TOTAL_FONT
        ws.cell(r_idx, 2).alignment = Alignment(horizontal="left")
        ws.cell(r_idx, 5).font = DATA_FONT
        c7 = ws.cell(r_idx, 7)
        c7.font = TOTAL_FONT
        c7.alignment = Alignment(horizontal="right")
        c7.number_format = '#,##0.00'

    # Signatures block removed as requested: "REMOVE PREPARED BY CHECKED BY AND AUTHOROZED SIGNS.. THIS PAGE DONT NEED SIGN"

    # PAGE BREAK: Page 1 prints Breakup Report only!
    breakup_last_row = ws.max_row
    ws.row_breaks.append(Break(id=breakup_last_row))

    det_headers = ['SL', 'Consumer ID', 'Payment Type', 'Amount', 'Payment Date', 'DOC NUMBER', 'ZONE', 'Phase']
    consumer_header_rows = []

    # ==========================================
    # SECTION 2: SINGLE PHASE (1PH) DETAILS (PAGE 2)
    # ==========================================
    phase1_records = [r for r in ag_records if r.get("phase") == 1]
    if phase1_records:
        ws.append([""] * total_cols)
        ws.row_dimensions[ws.max_row].height = 8

        # 1PH Main Section Title
        p1_title_row = ws.max_row + 1
        ws.append([f"1 PHASE (1PH) DETAILED CONSUMER LIST   |   AGENCY: {agency_name}"] + [""] * (total_cols - 1))
        ws.merge_cells(start_row=p1_title_row, start_column=1, end_row=p1_title_row, end_column=total_cols)
        ws.row_dimensions[p1_title_row].height = 24
        for c in range(1, total_cols + 1):
            cell = ws.cell(p1_title_row, c)
            cell.fill = BANNER_FILL
            cell.border = BORDER_THIN
        ws.cell(p1_title_row, 1).font = BANNER_FONT
        ws.cell(p1_title_row, 1).alignment = Alignment(horizontal="center", vertical="center")

        # 1PH Table Column Headers
        p1_hdr_row = ws.max_row + 1
        consumer_header_rows.append(p1_hdr_row)
        ws.append(det_headers)
        ws.row_dimensions[p1_hdr_row].height = 24
        for c in range(1, total_cols + 1):
            cell = ws.cell(p1_hdr_row, c)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORDER_THIN

        cur_part_idx = 1
        phase1_total_count = 0

        for ptype, part_label in PARTS_CONFIG_AGENCY:
            part_records = [
                r for r in phase1_records
                if str(r.get("payment_type") or "").strip().upper() == ptype
            ]
            if not part_records:
                continue

            phase1_total_count += len(part_records)

            # Part Header Banner (Dual Action removed, NO amount)
            pb_row = ws.max_row + 1
            part_title = f"PART {cur_part_idx}: 1 PHASE - {part_label}   |   COUNT: {len(part_records)}"
            ws.append([part_title] + [""] * (total_cols - 1))
            ws.merge_cells(start_row=pb_row, start_column=1, end_row=pb_row, end_column=total_cols)
            ws.row_dimensions[pb_row].height = 21
            for c in range(1, total_cols + 1):
                cell = ws.cell(pb_row, c)
                cell.fill = PART_FILL
                cell.border = BORDER_THIN
            ws.cell(pb_row, 1).font = PART_FONT
            ws.cell(pb_row, 1).alignment = Alignment(horizontal="left", vertical="center", indent=1)

            # Data rows (SL countable 1..N)
            for sl, r in enumerate(part_records, start=1):
                try:
                    amt_val = float(r["amount"])
                except (ValueError, TypeError):
                    amt_val = r["amount"]

                row_vals = [
                    sl, r["consumer_id"], r["payment_type"], amt_val,
                    r["payment_date"], r["doc_number"], r["zone"], r["phase"]
                ]
                ws.append(row_vals)
                cur_r = ws.max_row
                ws.row_dimensions[cur_r].height = 19
                for c in range(1, total_cols + 1):
                    cell = ws.cell(cur_r, c)
                    cell.font = DATA_FONT
                    cell.border = BORDER_THIN
                    if c == 4 and isinstance(amt_val, (int, float)):
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        cell.number_format = '#,##0.00'
                    else:
                        cell.alignment = Alignment(horizontal="center", vertical="center")

            # Part Details Subtotal Row: NO TOTAL AMOUNT, MERGED AND CENTERED TOTAL COUNT
            sub_row = ws.max_row + 1
            ws.append([f"TOTAL 1PH {ptype} RECORDS: {len(part_records)}"] + [""] * (total_cols - 1))
            ws.merge_cells(start_row=sub_row, start_column=1, end_row=sub_row, end_column=total_cols)
            ws.row_dimensions[sub_row].height = 20
            for c in range(1, total_cols + 1):
                cell = ws.cell(sub_row, c)
                cell.fill = SUBTOTAL_FILL
                cell.border = BORDER_DOUBLE_BOTTOM
            mid_c = ws.cell(sub_row, 1)
            mid_c.font = SUBTOTAL_FONT
            mid_c.alignment = Alignment(horizontal="center", vertical="center")

            # Spacing after parts
            ws.append([""] * total_cols)
            ws.row_dimensions[ws.max_row].height = 10

            cur_part_idx += 1

        # Total 1-Phase Row (Merged and Centered count, NO amount)
        p1_tot_r = ws.max_row + 1
        ws.append([f"TOTAL SINGLE PHASE (1PH) RECORDS: {phase1_total_count}"] + [""] * (total_cols - 1))
        ws.merge_cells(start_row=p1_tot_r, start_column=1, end_row=p1_tot_r, end_column=total_cols)
        ws.row_dimensions[p1_tot_r].height = 22
        for c in range(1, total_cols + 1):
            cell = ws.cell(p1_tot_r, c)
            cell.fill = TOTAL_FILL
            cell.border = BORDER_DOUBLE_BOTTOM
        ws.cell(p1_tot_r, 1).font = TOTAL_FONT
        ws.cell(p1_tot_r, 1).alignment = Alignment(horizontal="center", vertical="center")

    # ==========================================
    # PAGE BREAK: 1PH AND 3PH ON DIFFERENT PAGES!
    # ==========================================
    phase3_records = [r for r in ag_records if r.get("phase") == 3]
    if phase3_records:
        p1_last_row = ws.max_row
        ws.row_breaks.append(Break(id=p1_last_row))

        # ==========================================
        # SECTION 3: THREE PHASE (3PH) DETAILS
        # ==========================================
        ws.append([""] * total_cols)
        ws.row_dimensions[ws.max_row].height = 8

        # 3PH Main Section Title
        p3_title_row = ws.max_row + 1
        ws.append([f"3 PHASE (3PH) DETAILED CONSUMER LIST   |   AGENCY: {agency_name}"] + [""] * (total_cols - 1))
        ws.merge_cells(start_row=p3_title_row, start_column=1, end_row=p3_title_row, end_column=total_cols)
        ws.row_dimensions[p3_title_row].height = 24
        for c in range(1, total_cols + 1):
            cell = ws.cell(p3_title_row, c)
            cell.fill = BANNER_FILL
            cell.border = BORDER_THIN
        ws.cell(p3_title_row, 1).font = BANNER_FONT
        ws.cell(p3_title_row, 1).alignment = Alignment(horizontal="center", vertical="center")

        # 3PH Table Column Headers
        p3_hdr_row = ws.max_row + 1
        ws.append(det_headers)
        ws.row_dimensions[p3_hdr_row].height = 24
        for c in range(1, total_cols + 1):
            cell = ws.cell(p3_hdr_row, c)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORDER_THIN

        cur_part_idx = 1
        phase3_total_count = 0

        for ptype, part_label in PARTS_CONFIG_AGENCY:
            part_records = [
                r for r in phase3_records
                if str(r.get("payment_type") or "").strip().upper() == ptype
            ]
            if not part_records:
                continue

            phase3_total_count += len(part_records)

            # Part Header Banner (Dual Action removed, NO amount)
            pb_row = ws.max_row + 1
            part_title = f"PART {cur_part_idx}: 3 PHASE - {part_label}   |   COUNT: {len(part_records)}"
            ws.append([part_title] + [""] * (total_cols - 1))
            ws.merge_cells(start_row=pb_row, start_column=1, end_row=pb_row, end_column=total_cols)
            ws.row_dimensions[pb_row].height = 21
            for c in range(1, total_cols + 1):
                cell = ws.cell(pb_row, c)
                cell.fill = PART_FILL
                cell.border = BORDER_THIN
            ws.cell(pb_row, 1).font = PART_FONT
            ws.cell(pb_row, 1).alignment = Alignment(horizontal="left", vertical="center", indent=1)

            # Data rows (SL countable 1..N)
            for sl, r in enumerate(part_records, start=1):
                try:
                    amt_val = float(r["amount"])
                except (ValueError, TypeError):
                    amt_val = r["amount"]

                row_vals = [
                    sl, r["consumer_id"], r["payment_type"], amt_val,
                    r["payment_date"], r["doc_number"], r["zone"], r["phase"]
                ]
                ws.append(row_vals)
                cur_r = ws.max_row
                ws.row_dimensions[cur_r].height = 19
                for c in range(1, total_cols + 1):
                    cell = ws.cell(cur_r, c)
                    cell.font = DATA_FONT
                    cell.border = BORDER_THIN
                    if c == 4 and isinstance(amt_val, (int, float)):
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        cell.number_format = '#,##0.00'
                    else:
                        cell.alignment = Alignment(horizontal="center", vertical="center")

            # Part Details Subtotal Row: NO TOTAL AMOUNT, MERGED AND CENTERED TOTAL COUNT
            sub_row = ws.max_row + 1
            ws.append([f"TOTAL 3PH {ptype} RECORDS: {len(part_records)}"] + [""] * (total_cols - 1))
            ws.merge_cells(start_row=sub_row, start_column=1, end_row=sub_row, end_column=total_cols)
            ws.row_dimensions[sub_row].height = 20
            for c in range(1, total_cols + 1):
                cell = ws.cell(sub_row, c)
                cell.fill = SUBTOTAL_FILL
                cell.border = BORDER_DOUBLE_BOTTOM
            mid_c = ws.cell(sub_row, 1)
            mid_c.font = SUBTOTAL_FONT
            mid_c.alignment = Alignment(horizontal="center", vertical="center")

            # Spacing after parts
            ws.append([""] * total_cols)
            ws.row_dimensions[ws.max_row].height = 10

            cur_part_idx += 1

        # Total 3-Phase Row (Merged and Centered count, NO amount)
        p3_tot_r = ws.max_row + 1
        ws.append([f"TOTAL THREE PHASE (3PH) RECORDS: {phase3_total_count}"] + [""] * (total_cols - 1))
        ws.merge_cells(start_row=p3_tot_r, start_column=1, end_row=p3_tot_r, end_column=total_cols)
        ws.row_dimensions[p3_tot_r].height = 22
        for c in range(1, total_cols + 1):
            cell = ws.cell(p3_tot_r, c)
            cell.fill = TOTAL_FILL
            cell.border = BORDER_DOUBLE_BOTTOM
        ws.cell(p3_tot_r, 1).font = TOTAL_FONT
        ws.cell(p3_tot_r, 1).alignment = Alignment(horizontal="center", vertical="center")

    # OVERALL GRAND TOTAL COUNT ROW (Merged & Centered, NO amount)
    gtot_r = ws.max_row + 1
    ws.append([f"OVERALL TOTAL RECORDS: {m['grand_count']}"] + [""] * (total_cols - 1))
    ws.merge_cells(start_row=gtot_r, start_column=1, end_row=gtot_r, end_column=total_cols)
    ws.row_dimensions[gtot_r].height = 24
    for c in range(1, total_cols + 1):
        cell = ws.cell(gtot_r, c)
        cell.fill = TOTAL_FILL
        cell.border = BORDER_DOUBLE_BOTTOM
    ws.cell(gtot_r, 1).font = TOTAL_FONT
    ws.cell(gtot_r, 1).alignment = Alignment(horizontal="center", vertical="center")

    # Repeat consumer table headers across subsequent pages
    repeat_hdr = consumer_header_rows[0] if consumer_header_rows else ""
    title_rows_str = f"{repeat_hdr}:{repeat_hdr}" if repeat_hdr else ""

    _setup_portrait_page(ws, title_rows=title_rows_str, fit_to_height=0)
    _apply_value_based_widths(ws, det_headers, data_start_row=3)
    return ws


def _create_raw_data_sheet(wb, headers, records, is_master=False):
    """
    Creates a pure, completely unformatted data sheet:
    Just column headers on Row 1 and pure raw records below.
    Zero formatting, zero colors, zero custom fonts.
    """
    ws = wb.create_sheet(title="RAW_DATA")
    ws.append(headers)
    for sl, r in enumerate(records, start=1):
        if is_master:
            ws.append([
                sl, r["consumer_id"], r["payment_type"], r["amount"],
                r["payment_date"], r["doc_number"], r["zone"],
                r["agency"], r["phase"]
            ])
        else:
            ws.append([
                sl, r["consumer_id"], r["payment_type"], r["amount"],
                r["payment_date"], r["doc_number"], r["zone"],
                r["phase"]
            ])
    return ws


def export_agency_workbooks(processed_result, output_dir):
    """
    Exports individual <Agency>.xlsx workbooks for each agency.
    Sheet structure:
    1. PRINT_ALL: The all-in-one printable bill:
       - Page 1: Breakup report (no signatures, no dual action word)
       - Page Break
       - Page 2: 1PH details (DR, DC, RC with countable SL and merged & centered total counts; NO amount; spacings)
       - Page Break (1PH and 3PH on different pages)
       - Page 3: 3PH details (DR, DC, RC with countable SL and merged & centered total counts; NO amount; spacings)
    2. BREAKUP REPORT: Standalone 1-page Breakup summary.
    3. Phase 1: Standalone Phase 1 detailed list with DR, DC, RC.
    4. Phase 3: Standalone Phase 3 detailed list with DR, DC, RC.
    5. RAW_DATA: Pure unformatted raw table data.
    """
    os.makedirs(output_dir, exist_ok=True)
    records = processed_result.get("records", [])
    date_range_str = _get_date_range_str(processed_result)

    agencies_map = {}
    for r in records:
        ag = r.get("agency") or "UNASSIGNED"
        if ag not in agencies_map:
            agencies_map[ag] = []
        agencies_map[ag].append(r)

    exported_files = []
    headers = ['SL', 'Consumer ID', 'Payment Type', 'Amount', 'Payment Date', 'DOC NUMBER', 'ZONE', 'Phase']

    sorted_agencies = sorted([a for a in agencies_map if a != "UNASSIGNED"])
    if "UNASSIGNED" in agencies_map:
        sorted_agencies.append("UNASSIGNED")

    for file_num, agency_name in enumerate(sorted_agencies, start=3):
        ag_records = agencies_map[agency_name]
        wb = openpyxl.Workbook()
        default_sheet = wb.active

        # Sheet 1: The All-In-One Printable Bill
        _create_agency_print_all_sheet(wb, agency_name, processed_result, ag_records)

        # Sheet 2: Standalone Breakup Report
        _create_agency_breakup_sheet(wb, agency_name, processed_result, ag_records)

        # Sheets 3 & 4: Standalone Phase 1 and Phase 3 sheets
        for phase_num in (1, 3):
            ws = wb.create_sheet(title=f"Phase {phase_num}")

            banner_text = f"AGENCY: {agency_name}   |   PHASE {phase_num}   |   DATE RANGE: {date_range_str}"
            ws.append([banner_text])
            _style_banner_row(ws, 1, len(headers), banner_text, fill=BANNER_FILL, font=BANNER_FONT, height=24)

            ws.append(headers)
            ws.row_dimensions[2].height = 24
            for col_idx, h in enumerate(headers, start=1):
                cell = ws.cell(2, col_idx)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = BORDER_THIN

            phase_records = [r for r in ag_records if r.get("phase") == phase_num]
            cur_part_idx = 1
            phase_total_count = 0

            for ptype, part_label in PARTS_CONFIG_AGENCY:
                part_records = [
                    r for r in phase_records
                    if str(r.get("payment_type") or "").strip().upper() == ptype
                ]
                if not part_records:
                    continue

                phase_total_count += len(part_records)

                # Part Banner (NO amount, Dual Action removed)
                banner_row = ws.max_row + 1
                part_title = f"PART {cur_part_idx}: {phase_num} PHASE - {part_label}   |   COUNT: {len(part_records)}"
                ws.append([part_title])
                ws.merge_cells(start_row=banner_row, start_column=1, end_row=banner_row, end_column=len(headers))
                ws.row_dimensions[banner_row].height = 21
                for c in range(1, len(headers) + 1):
                    cell = ws.cell(banner_row, c)
                    cell.fill = PART_FILL
                    cell.border = BORDER_THIN
                top_c = ws.cell(banner_row, 1)
                top_c.font = PART_FONT
                top_c.alignment = Alignment(horizontal="left", vertical="center", indent=1)

                for sl, r in enumerate(part_records, start=1):
                    try:
                        amt_val = float(r["amount"])
                    except (ValueError, TypeError):
                        amt_val = r["amount"]

                    row_vals = [
                        sl, r["consumer_id"], r["payment_type"], amt_val,
                        r["payment_date"], r["doc_number"], r["zone"], r["phase"]
                    ]
                    ws.append(row_vals)
                    cur_row = ws.max_row
                    ws.row_dimensions[cur_row].height = 19

                    for c in range(1, len(row_vals) + 1):
                        cell = ws.cell(cur_row, c)
                        cell.font = DATA_FONT
                        cell.border = BORDER_THIN
                        if c == 4 and isinstance(amt_val, (int, float)):
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                            cell.number_format = '#,##0.00'
                        else:
                            cell.alignment = Alignment(horizontal="center", vertical="center")

                # Subtotal row: NO TOTAL AMOUNT, MERGED AND CENTERED TOTAL COUNT
                subtotal_row = ws.max_row + 1
                ws.append([f"TOTAL {phase_num}PH {ptype} RECORDS: {len(part_records)}"] + [""] * (len(headers) - 1))
                ws.merge_cells(start_row=subtotal_row, start_column=1, end_row=subtotal_row, end_column=len(headers))
                ws.row_dimensions[subtotal_row].height = 20
                for c in range(1, len(headers) + 1):
                    cell = ws.cell(subtotal_row, c)
                    cell.fill = SUBTOTAL_FILL
                    cell.border = BORDER_DOUBLE_BOTTOM
                mid_c = ws.cell(subtotal_row, 1)
                mid_c.font = SUBTOTAL_FONT
                mid_c.alignment = Alignment(horizontal="center", vertical="center")

                # Spacing after parts
                ws.append([""] * len(headers))
                ws.row_dimensions[ws.max_row].height = 10

                cur_part_idx += 1

            if phase_total_count > 0:
                tot_row = ws.max_row + 1
                ws.append([f"TOTAL {phase_num} PHASE ({phase_num}PH) RECORDS: {phase_total_count}"] + [""] * (len(headers) - 1))
                ws.merge_cells(start_row=tot_row, start_column=1, end_row=tot_row, end_column=len(headers))
                ws.row_dimensions[tot_row].height = 22
                for c in range(1, len(headers) + 1):
                    cell = ws.cell(tot_row, c)
                    cell.fill = TOTAL_FILL
                    cell.border = BORDER_DOUBLE_BOTTOM
                ws.cell(tot_row, 1).font = TOTAL_FONT
                ws.cell(tot_row, 1).alignment = Alignment(horizontal="center", vertical="center")
            else:
                empty_row = ws.max_row + 1
                ws.append(["", f"No Phase {phase_num} records for {agency_name}", "", "", "", "", "", ""])
                ws.merge_cells(start_row=empty_row, start_column=1, end_row=empty_row, end_column=len(headers))
                for c in range(1, len(headers) + 1):
                    cell = ws.cell(empty_row, c)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.font = DATA_FONT

            _setup_portrait_page(ws, title_rows='1:2', fit_to_height=0)
            _apply_value_based_widths(ws, headers, data_start_row=3)

        # Sheet 5: Pure Raw Data Sheet (Zero formatting)
        _create_raw_data_sheet(wb, headers, ag_records, is_master=False)

        wb.remove(default_sheet)

        safe_name = "".join(c for c in agency_name if c.isalnum() or c in (' ', '_', '-')).strip()
        file_path = os.path.join(output_dir, f"{file_num:02d}_{safe_name}.xlsx")
        wb.save(file_path)
        exported_files.append(file_path)

    return exported_files


def export_summary_report(processed_result, output_dir):
    """
    Exports the PO Summary Report matching the layout of the REPORT sheet in LIST SOFTWARE V2.xlsm.
    Prefixed with 01_ for sorting. Formatted with light headers and horizontal centering.
    """
    os.makedirs(output_dir, exist_ok=True)
    file_path = os.path.join(output_dir, "01_PO_Summary_Report.xlsx")
    date_range_str = _get_date_range_str(processed_result)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "REPORT"

    total_cols = 18

    # Row 1: Banner
    banner_text = f"PO SUMMARY BILLING REPORT   |   DATE RANGE: {date_range_str}"
    ws.append([banner_text])
    _style_banner_row(ws, 1, total_cols, banner_text, fill=BANNER_FILL, font=BANNER_FONT, height=24)

    # Row 2: Group Headers
    ws.append([""] * total_cols)
    ws.row_dimensions[2].height = 22
    ws.cell(2, 1, "Agency")
    ws.cell(2, 2, "1 PHASE (1PH)")
    ws.cell(2, 8, "3 PHASE (3PH)")
    ws.cell(2, 14, "TOTAL PAYABLE")

    ws.merge_cells("B2:G2")
    ws.merge_cells("H2:M2")
    ws.merge_cells("N2:R2")

    for c in range(1, total_cols + 1):
        cell = ws.cell(2, c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER_THIN

    # Row 3: Sub-headers
    sub_headers = [
        "Agency",
        "DC", "Rate", "RC", "Rate", "DR", "Rate",
        "DC", "Rate", "RC", "Rate", "DR", "Rate",
        "Total Amount", "Total DC", "Total RC", "Total DR", "Total Records"
    ]
    ws.append(sub_headers)
    ws.row_dimensions[3].height = 22

    for c in range(1, total_cols + 1):
        cell = ws.cell(3, c)
        cell.font = SUBHEADER_FONT
        cell.fill = SUBHEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_THIN

    # Data Rows (Row 4+)
    summary_rows = processed_result.get("summary", [])
    for item in summary_rows:
        row_vals = [
            item["agency"],
            item["c_1ph_dc"], item["r_1ph_dc"],
            item["c_1ph_rc"], item["r_1ph_rc"],
            item["c_1ph_dr"], item["r_1ph_dr"],
            item["c_3ph_dc"], item["r_3ph_dc"],
            item["c_3ph_rc"], item["r_3ph_rc"],
            item["c_3ph_dr"], item["r_3ph_dr"],
            item["total_amount"],
            item["total_dc"], item["total_rc"], item["total_dr"], item["total_count"]
        ]
        ws.append(row_vals)
        cur_row = ws.max_row
        ws.row_dimensions[cur_row].height = 19
        for c in range(1, total_cols + 1):
            cell = ws.cell(cur_row, c)
            cell.font = DATA_FONT
            cell.border = BORDER_THIN
            if c == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if c == 14:
                    cell.number_format = '#,##0.00'

    # Grand Total Row
    totals = processed_result.get("totals", {})
    total_row = [
        "TOTAL",
        totals.get("1ph_dc", 0), "",
        totals.get("1ph_rc", 0), "",
        totals.get("1ph_dr", 0), "",
        totals.get("3ph_dc", 0), "",
        totals.get("3ph_rc", 0), "",
        totals.get("3ph_dr", 0), "",
        totals.get("total_amount", 0.0),
        totals.get("1ph_dc", 0) + totals.get("3ph_dc", 0),
        totals.get("1ph_rc", 0) + totals.get("3ph_rc", 0),
        totals.get("1ph_dr", 0) + totals.get("3ph_dr", 0),
        totals.get("total_records", 0)
    ]
    ws.append(total_row)
    cur_row = ws.max_row
    ws.row_dimensions[cur_row].height = 22
    for c in range(1, total_cols + 1):
        cell = ws.cell(cur_row, c)
        cell.font = TOTAL_FONT
        cell.fill = TOTAL_FILL
        cell.border = BORDER_THIN
        if c == 1:
            cell.alignment = Alignment(horizontal="left", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            if c == 14:
                cell.number_format = '#,##0.00'

    # Phase summary block below
    ws.cell(cur_row + 2, 2, "PHASE SUMMARY").font = TOTAL_FONT
    ws.cell(cur_row + 3, 2, "1PH Total:").font = TOTAL_FONT
    ws.cell(cur_row + 3, 3, totals.get("1ph_total", 0)).font = TOTAL_FONT
    ws.cell(cur_row + 4, 2, "3PH Total:").font = TOTAL_FONT
    ws.cell(cur_row + 4, 3, totals.get("3ph_total", 0)).font = TOTAL_FONT
    ws.cell(cur_row + 5, 2, "Overall Total Records:").font = TOTAL_FONT
    ws.cell(cur_row + 5, 3, totals.get("total_records", 0)).font = TOTAL_FONT
    ws.cell(cur_row + 6, 2, "Overall P.O. Amount (Rs):").font = TOTAL_FONT
    tot_amt_cell = ws.cell(cur_row + 6, 3, totals.get("total_amount", 0.0))
    tot_amt_cell.font = TOTAL_FONT
    tot_amt_cell.number_format = '#,##0.00'

    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.4
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:3'
    ws.oddFooter.right.text = "Page &P of &N"
    ws.evenFooter.right.text = "Page &P of &N"
    ws.oddFooter.left.text = ""
    ws.evenFooter.left.text = ""
    ws.print_options.gridLines = True

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        if col[0].column == 1:
            ws.column_dimensions[col_letter].width = 20
        elif col[0].column == 14:
            ws.column_dimensions[col_letter].width = 14
        else:
            ws.column_dimensions[col_letter].width = 9

    wb.save(file_path)
    return file_path


def export_master_list(processed_result, output_dir):
    """
    Exports the complete MAIN LIST containing all merged records.
    - Sheet 1 (MAIN LIST): Continuous unparted list with continuous SL (1..N), light headers, and merged total count.
    - Sheet 2 (RAW_DATA): Pure unformatted raw list (zero styling/colors).
    """
    os.makedirs(output_dir, exist_ok=True)
    file_path = os.path.join(output_dir, "02_MAIN_LIST.xlsx")
    date_range_str = _get_date_range_str(processed_result)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MAIN LIST"

    headers = ['SL', 'Consumer ID', 'Payment Type', 'Amount', 'Payment Date', 'DOC NUMBER', 'ZONE', 'Agency', 'Phase']

    # Row 1: Light Banner
    banner_text = f"DCRC RECONCILIATION MAIN LIST   |   DATE RANGE: {date_range_str}"
    ws.append([banner_text])
    _style_banner_row(ws, 1, len(headers), banner_text, fill=BANNER_FILL, font=BANNER_FONT, height=24)

    # Row 2: Light Column Headers
    ws.append(headers)
    ws.row_dimensions[2].height = 24
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(2, col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_THIN

    all_records = processed_result.get("records", [])

    # Continuous unparted rows with continuous SL from 1 to N
    for sl, r in enumerate(all_records, start=1):
        try:
            amt_val = float(r["amount"])
        except (ValueError, TypeError):
            amt_val = r["amount"]

        row_vals = [
            sl,
            r["consumer_id"],
            r["payment_type"],
            amt_val,
            r["payment_date"],
            r["doc_number"],
            r["zone"],
            r["agency"],
            r["phase"]
        ]
        ws.append(row_vals)
        cur_row = ws.max_row
        ws.row_dimensions[cur_row].height = 19

        for c in range(1, len(row_vals) + 1):
            cell = ws.cell(cur_row, c)
            cell.font = DATA_FONT
            cell.border = BORDER_THIN
            if c == 4 and isinstance(amt_val, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0.00'
            elif c == 8:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # Bottom Total Row: Merged and centered total count (NO amount)
    if all_records:
        gt_row = ws.max_row + 1
        ws.append([f"TOTAL CONSUMER RECORDS: {len(all_records)}"] + [""] * (len(headers) - 1))
        ws.merge_cells(start_row=gt_row, start_column=1, end_row=gt_row, end_column=len(headers))
        ws.row_dimensions[gt_row].height = 24
        for c in range(1, len(headers) + 1):
            cell = ws.cell(gt_row, c)
            cell.fill = TOTAL_FILL
            cell.border = BORDER_DOUBLE_BOTTOM
        ws.cell(gt_row, 1).font = TOTAL_FONT
        ws.cell(gt_row, 1).alignment = Alignment(horizontal="center", vertical="center")

    # Configure Portrait A4 with repeating headers and horizontal centering
    _setup_portrait_page(ws, title_rows='1:2', fit_to_height=0)
    _apply_value_based_widths(ws, headers, data_start_row=3)

    # Sheet 2: Pure Raw List (Zero formatting)
    _create_raw_data_sheet(wb, headers, all_records, is_master=True)

    wb.save(file_path)
    return file_path
