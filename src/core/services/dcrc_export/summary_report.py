import os
import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from .styles import (
    BANNER_FILL, BANNER_FONT, HEADER_FILL, HEADER_FONT,
    SUBHEADER_FILL, SUBHEADER_FONT, DATA_FONT,
    TOTAL_FILL, TOTAL_FONT, BORDER_THIN, _get_date_range_str,
    _style_banner_row
)


def export_summary_report(processed_result, output_dir):
    """
    Exports the PO Summary Report matching the layout of the REPORT sheet in LIST SOFTWARE V2.xlsm.
    Prefixed with 01_ for sorting. Formatted with light headers and horizontal centering.
    """
    os.makedirs(output_dir, exist_ok=True)
    file_path = os.path.join(output_dir, "01_PO_Summary_Report.xlsx")
    date_range_str = _get_date_range_str(processed_result)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "REPORT"

    total_cols = 18

    # Row 1: Banner
    banner_text = f"PO SUMMARY BILLING REPORT   |   DATE RANGE: {date_range_str}"
    ws.append([banner_text])
    _style_banner_row(ws, 1, total_cols, banner_text, fill=BANNER_FILL, font=BANNER_FONT, height=24)

    # Row 2: Group Headers
    ws.append([""] * total_cols)
    ws.row_dimensions[2].height = 22
    ws.cell(2, 1, "Agency")
    ws.cell(2, 2, "1 PHASE (1PH)")
    ws.cell(2, 8, "3 PHASE (3PH)")
    ws.cell(2, 14, "TOTAL PAYABLE")

    ws.merge_cells("B2:G2")
    ws.merge_cells("H2:M2")
    ws.merge_cells("N2:R2")

    for c in range(1, total_cols + 1):
        cell = ws.cell(2, c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER_THIN

    # Row 3: Sub-headers
    sub_headers = [
        "Agency",
        "DC", "Rate", "RC", "Rate", "DR", "Rate",
        "DC", "Rate", "RC", "Rate", "DR", "Rate",
        "Total Amount", "Total DC", "Total RC", "Total DR", "Total Records"
    ]
    ws.append(sub_headers)
    ws.row_dimensions[3].height = 22

    for c in range(1, total_cols + 1):
        cell = ws.cell(3, c)
        cell.font = SUBHEADER_FONT
        cell.fill = SUBHEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_THIN

    # Data Rows (Row 4+)
    summary_rows = processed_result.get("summary", [])
    for item in summary_rows:
        row_vals = [
            item["agency"],
            item["c_1ph_dc"], item["r_1ph_dc"],
            item["c_1ph_rc"], item["r_1ph_rc"],
            item["c_1ph_dr"], item["r_1ph_dr"],
            item["c_3ph_dc"], item["r_3ph_dc"],
            item["c_3ph_rc"], item["r_3ph_rc"],
            item["c_3ph_dr"], item["r_3ph_dr"],
            item["total_amount"],
            item["total_dc"], item["total_rc"], item["total_dr"], item["total_count"]
        ]
        ws.append(row_vals)
        cur_row = ws.max_row
        ws.row_dimensions[cur_row].height = 19
        for c in range(1, total_cols + 1):
            cell = ws.cell(cur_row, c)
            cell.font = DATA_FONT
            cell.border = BORDER_THIN
            if c == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if c == 14:
                    cell.number_format = '#,##0.00'

    # Grand Total Row
    totals = processed_result.get("totals", {})
    total_row = [
        "TOTAL",
        totals.get("1ph_dc", 0), "",
        totals.get("1ph_rc", 0), "",
        totals.get("1ph_dr", 0), "",
        totals.get("3ph_dc", 0), "",
        totals.get("3ph_rc", 0), "",
        totals.get("3ph_dr", 0), "",
        totals.get("total_amount", 0.0),
        totals.get("1ph_dc", 0) + totals.get("3ph_dc", 0),
        totals.get("1ph_rc", 0) + totals.get("3ph_rc", 0),
        totals.get("1ph_dr", 0) + totals.get("3ph_dr", 0),
        totals.get("total_records", 0)
    ]
    ws.append(total_row)
    cur_row = ws.max_row
    ws.row_dimensions[cur_row].height = 22
    for c in range(1, total_cols + 1):
        cell = ws.cell(cur_row, c)
        cell.font = TOTAL_FONT
        cell.fill = TOTAL_FILL
        cell.border = BORDER_THIN
        if c == 1:
            cell.alignment = Alignment(horizontal="left", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            if c == 14:
                cell.number_format = '#,##0.00'

    # Phase summary block below
    ws.cell(cur_row + 2, 2, "PHASE SUMMARY").font = TOTAL_FONT
    ws.cell(cur_row + 3, 2, "1PH Total:").font = TOTAL_FONT
    ws.cell(cur_row + 3, 3, totals.get("1ph_total", 0)).font = TOTAL_FONT
    ws.cell(cur_row + 4, 2, "3PH Total:").font = TOTAL_FONT
    ws.cell(cur_row + 4, 3, totals.get("3ph_total", 0)).font = TOTAL_FONT
    ws.cell(cur_row + 5, 2, "Overall Total Records:").font = TOTAL_FONT
    ws.cell(cur_row + 5, 3, totals.get("total_records", 0)).font = TOTAL_FONT
    ws.cell(cur_row + 6, 2, "Overall P.O. Amount (Rs):").font = TOTAL_FONT
    tot_amt_cell = ws.cell(cur_row + 6, 3, totals.get("total_amount", 0.0))
    tot_amt_cell.font = TOTAL_FONT
    tot_amt_cell.number_format = '#,##0.00'

    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.4
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:3'
    ws.oddFooter.right.text = "Page &P of &N"
    ws.evenFooter.right.text = "Page &P of &N"
    ws.oddFooter.left.text = ""
    ws.evenFooter.left.text = ""
    ws.print_options.gridLines = True

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        if col[0].column == 1:
            ws.column_dimensions[col_letter].width = 20
        elif col[0].column == 14:
            ws.column_dimensions[col_letter].width = 14
        else:
            ws.column_dimensions[col_letter].width = 9

    wb.save(file_path)
    return file_path
