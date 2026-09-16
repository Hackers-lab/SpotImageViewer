import os
import openpyxl
from openpyxl.styles import Alignment

from .styles import (
    BANNER_FILL, BANNER_FONT, HEADER_FILL, HEADER_FONT,
    DATA_FONT, TOTAL_FILL, TOTAL_FONT, BORDER_THIN,
    BORDER_DOUBLE_BOTTOM, _setup_portrait_page,
    _get_date_range_str, _apply_value_based_widths,
    _style_banner_row
)
from .agency_sheets import _create_raw_data_sheet


def export_master_list(processed_result, output_dir):
    """
    Exports the complete MAIN LIST containing all merged records.
    - Sheet 1 (MAIN LIST): Continuous unparted list with continuous SL (1..N), light headers, and merged total count.
    - Sheet 2 (RAW_DATA): Pure unformatted raw list (zero styling/colors).
    """
    os.makedirs(output_dir, exist_ok=True)
    file_path = os.path.join(output_dir, "02_MAIN_LIST.xlsx")
    date_range_str = _get_date_range_str(processed_result)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MAIN LIST"

    headers = ['SL', 'Consumer ID', 'Payment Type', 'Amount', 'Payment Date', 'DOC NUMBER', 'ZONE', 'Agency', 'Phase']

    # Row 1: Light Banner
    banner_text = f"DCRC RECONCILIATION MAIN LIST   |   DATE RANGE: {date_range_str}"
    ws.append([banner_text])
    _style_banner_row(ws, 1, len(headers), banner_text, fill=BANNER_FILL, font=BANNER_FONT, height=24)

    # Row 2: Light Column Headers
    ws.append(headers)
    ws.row_dimensions[2].height = 24
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(2, col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_THIN

    all_records = processed_result.get("records", [])

    # Continuous unparted rows with continuous SL from 1 to N
    for sl, r in enumerate(all_records, start=1):
        try:
            amt_val = float(r["amount"])
        except (ValueError, TypeError):
            amt_val = r["amount"]

        row_vals = [
            sl,
            r["consumer_id"],
            r["payment_type"],
            amt_val,
            r["payment_date"],
            r["doc_number"],
            r["zone"],
            r["agency"],
            r["phase"]
        ]
        ws.append(row_vals)
        cur_row = ws.max_row
        ws.row_dimensions[cur_row].height = 19

        for c in range(1, len(row_vals) + 1):
            cell = ws.cell(cur_row, c)
            cell.font = DATA_FONT
            cell.border = BORDER_THIN
            if c == 4 and isinstance(amt_val, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0.00'
            elif c == 8:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # Bottom Total Row: Merged and centered total count (NO amount)
    if all_records:
        gt_row = ws.max_row + 1
        ws.append([f"TOTAL CONSUMER RECORDS: {len(all_records)}"] + [""] * (len(headers) - 1))
        ws.merge_cells(start_row=gt_row, start_column=1, end_row=gt_row, end_column=len(headers))
        ws.row_dimensions[gt_row].height = 24
        for c in range(1, len(headers) + 1):
            cell = ws.cell(gt_row, c)
            cell.fill = TOTAL_FILL
            cell.border = BORDER_DOUBLE_BOTTOM
        ws.cell(gt_row, 1).font = TOTAL_FONT
        ws.cell(gt_row, 1).alignment = Alignment(horizontal="center", vertical="center")

    # Configure Portrait A4 with repeating headers and horizontal centering
    _setup_portrait_page(ws, title_rows='1:2', fit_to_height=0)
    _apply_value_based_widths(ws, headers, data_start_row=3)

    # Sheet 2: Pure Raw List (Zero formatting)
    _create_raw_data_sheet(wb, headers, all_records, is_master=True)

    wb.save(file_path)
    return file_path
